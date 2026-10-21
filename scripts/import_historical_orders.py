from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5
from zoneinfo import ZoneInfo

import openpyxl
from sqlalchemy import func, text
from sqlmodel import Session, select

from gastroflow.data.database import engine
from gastroflow.domain.enums import EstadoPedido, FormaPago, UnidadVenta
from gastroflow.models import (
    Agregado,
    Categoria,
    Cliente,
    Pedido,
    PedidoItem,
    PedidoItemAgregado,
    PedidoPago,
    Producto,
    ZonaEnvio,
)


EXPECTED_HEADERS = (
    "ID", "Fecha", "Pedido", "Categoria", "Persona", "Delivery",
    "Pago", "Monto", "Estado", "Envio", "Observaciones",
)
EXPECTED_SOURCE_ROWS = 663
EXPECTED_ORDER_COUNT = 242
EXPECTED_FIRST_DATE = date(2026, 6, 13)
EXPECTED_LAST_DATE = date(2026, 10, 10)
CUTOFF_DATE = date(2026, 7, 23)
MONEY_QUANT = Decimal("0.01")
LOCAL_TZ = ZoneInfo("America/Argentina/Buenos_Aires")

CATEGORY_CODES = {
    "lista para hornear": "LISTA_HORNO",
    "lista para comer": "LISTA_COMER",
    "pre pizzas": "PRE_PIZZAS",
    "pan arabe": "PAN_ARABE",
    "mayorista": "MAYORISTA",
}

SINGLE_PRODUCTS = {
    "muzza": "MUZZA",
    "especial": "ESPECIAL",
    "calabresa": "CALABRESA",
    "fugazzeta": "FUGAZZETTA",
    "napolitana": "NAPOLITANA",
    "ternera": "TERNERA_HISTORICA",
}

COMBO_PRODUCTS = {
    "especial y calabresa": ("ESPECIAL", "CALABRESA"),
    "muzza y fuga": ("MUZZA", "FUGAZZETTA"),
    "calabresa y fuga": ("CALABRESA", "FUGAZZETTA"),
    "especial y fuga": ("ESPECIAL", "FUGAZZETTA"),
    "especial y napo": ("ESPECIAL", "NAPOLITANA"),
}

NAME_CORRECTIONS = {
    (date(2026, 6, 17), "futbol tdc"): "Bernardo Graffigna",
    (date(2026, 7, 1), "futbol tdc"): "Francisco Orlandi",
}


@dataclass(frozen=True)
class SourceRow:
    excel_row: int
    delivery_date: date
    product_text: str
    category_text: str
    customer_name: str
    address: str
    payment_text: str
    line_total: Decimal
    state_text: str
    shipping: Decimal | None
    observation: str


def normalize_text(value: object) -> str:
    text_value = " ".join(str(value or "").strip().split())
    return "".join(
        character
        for character in unicodedata.normalize("NFKD", text_value)
        if not unicodedata.combining(character)
    ).casefold()


def money(value: object) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0.00")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(MONEY_QUANT)
    cleaned = str(value).strip().replace("$", "").replace(" ", "")
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    return Decimal(cleaned).quantize(MONEY_QUANT)


def load_source(path: Path) -> list[SourceRow]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Pedidos"]
    headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if headers != EXPECTED_HEADERS:
        raise RuntimeError(f"Encabezados inesperados: {headers!r}")

    rows: list[SourceRow] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if values[1] is None:
            continue
        raw_date = values[1]
        if not isinstance(raw_date, datetime):
            raise RuntimeError(f"Fecha inválida en fila {excel_row}: {raw_date!r}")
        rows.append(SourceRow(
            excel_row=excel_row,
            delivery_date=raw_date.date(),
            product_text=str(values[2] or "").strip(),
            category_text=str(values[3] or "").strip(),
            customer_name=str(values[4] or "").strip(),
            address=str(values[5] or "").strip(),
            payment_text=str(values[6] or "").strip(),
            line_total=money(values[7]),
            state_text=str(values[8] or "").strip(),
            shipping=None if values[9] is None else money(values[9]),
            observation=str(values[10] or "").strip(),
        ))
    workbook.close()

    if len(rows) != EXPECTED_SOURCE_ROWS:
        raise RuntimeError(f"Se esperaban {EXPECTED_SOURCE_ROWS} filas y se encontraron {len(rows)}.")
    if min(row.delivery_date for row in rows) != EXPECTED_FIRST_DATE:
        raise RuntimeError("La primera fecha del archivo no coincide con la validada.")
    if max(row.delivery_date for row in rows) != EXPECTED_LAST_DATE:
        raise RuntimeError("La última fecha del archivo no coincide con la validada.")
    if {normalize_text(row.state_text) for row in rows} != {"entregado"}:
        raise RuntimeError("Todos los pedidos históricos deben estar entregados.")
    if {normalize_text(row.payment_text) for row in rows} - {"efectivo", "transferencia"}:
        raise RuntimeError("Se encontró una forma de pago no contemplada.")
    if any(row.line_total <= 0 for row in rows):
        raise RuntimeError("Todos los importes de línea deben ser mayores a cero.")
    return rows


def corrected_name(row: SourceRow) -> str:
    return NAME_CORRECTIONS.get(
        (row.delivery_date, normalize_text(row.customer_name)),
        " ".join(row.customer_name.split()),
    )


def group_key(row: SourceRow) -> tuple[object, ...]:
    name = normalize_text(corrected_name(row))
    if row.delivery_date == date(2026, 9, 17) and name == "daniela diaz":
        return row.delivery_date, name, normalize_text(row.address)
    return row.delivery_date, name


def group_source(rows: list[SourceRow]) -> list[list[SourceRow]]:
    grouped: dict[tuple[object, ...], list[SourceRow]] = {}
    for row in rows:
        grouped.setdefault(group_key(row), []).append(row)
    groups = sorted(grouped.values(), key=lambda group: (group[0].delivery_date, group[0].excel_row))
    if len(groups) != EXPECTED_ORDER_COUNT:
        raise RuntimeError(f"Se esperaban {EXPECTED_ORDER_COUNT} pedidos y se reconstruyeron {len(groups)}.")
    for group in groups:
        if len({normalize_text(row.address) for row in group}) != 1:
            raise RuntimeError(f"El pedido de la fila {group[0].excel_row} tiene más de una dirección.")
        if sum(row.shipping is not None for row in group) > 1:
            raise RuntimeError(f"El pedido de la fila {group[0].excel_row} tiene más de un costo de envío.")
    return groups


def explicit_zone(address: str) -> str | None:
    normalized = normalize_text(address)
    if re.search(r"(?:^|\s|-)smt(?:$|\s)", normalized):
        return "San Miguel de Tucumán"
    if re.search(r"(?:^|\s|-)yb(?:$|\s)", normalized) or "yerba buena" in normalized:
        return "Yerba Buena"
    if re.search(r"(?:^|\s|-)ald(?:$|\s)", normalized) or "alderetes" in normalized:
        return "Alderetes"
    return None


def normalized_address_key(address: str) -> str:
    normalized = normalize_text(address)
    return re.sub(r"\s*-\s*(?:smt|yb|ald)\s*$", "", normalized).strip()


def tariff_zone(row: SourceRow) -> str | None:
    if row.shipping is None:
        return None
    if row.delivery_date < CUTOFF_DATE:
        return {
            Decimal("1000.00"): "San Miguel de Tucumán",
            Decimal("1500.00"): "Yerba Buena",
        }.get(row.shipping)
    return {
        Decimal("500.00"): "Alderetes",
        Decimal("1500.00"): "San Miguel de Tucumán",
        Decimal("2000.00"): "Yerba Buena",
    }.get(row.shipping)


def address_zone_map(rows: list[SourceRow]) -> dict[str, str]:
    candidates: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if normalize_text(row.address) == "retiro en el local":
            continue
        zone = explicit_zone(row.address) or tariff_zone(row)
        if zone:
            candidates[normalized_address_key(row.address)].add(zone)
    return {
        address: next(iter(zones))
        for address, zones in candidates.items()
        if len(zones) == 1
    }


def resolve_group_zone(group: list[SourceRow], learned: dict[str, str]) -> str | None:
    address = group[0].address
    if normalize_text(address) == "retiro en el local":
        return None
    candidates = {
        zone
        for row in group
        if (zone := explicit_zone(row.address) or tariff_zone(row)) is not None
    }
    learned_zone = learned.get(normalized_address_key(address))
    if learned_zone:
        candidates.add(learned_zone)
    if not candidates and group[0].delivery_date < CUTOFF_DATE:
        # Under the original tariff, an unmarked free delivery corresponds to Alderetes.
        candidates.add("Alderetes")
    if len(candidates) > 1:
        raise RuntimeError(
            f"Zona contradictoria en pedido de fila {group[0].excel_row}: {sorted(candidates)}"
        )
    return next(iter(candidates)) if candidates else None


def payment_type(value: str) -> FormaPago:
    return FormaPago.EFECTIVO if normalize_text(value) == "efectivo" else FormaPago.TRANSFERENCIA


def parse_product(row: SourceRow) -> tuple[str, str | None, int]:
    normalized = normalize_text(row.product_text)
    if normalized in COMBO_PRODUCTS:
        first, second = COMBO_PRODUCTS[normalized]
        return first, second, 1
    if normalized.startswith("pre pizzas x"):
        quantity = int(normalized.rsplit("x", 1)[1])
        return "PREPIZZA", None, quantity
    if normalized.startswith("pan arabe x"):
        quantity = int(normalized.rsplit("x", 1)[1])
        return "PANARABE", None, quantity
    try:
        return SINGLE_PRODUCTS[normalized], None, 1
    except KeyError as exc:
        raise RuntimeError(f"Producto desconocido en fila {row.excel_row}: {row.product_text}") from exc


def group_observation(group: list[SourceRow]) -> str | None:
    observations = list(dict.fromkeys(row.observation for row in group if row.observation))
    is_pickup = normalize_text(group[0].address) == "retiro en el local"
    if not is_pickup and all(row.shipping is None for row in group):
        observations.append("Envío sin costo")
    return " | ".join(observations) or None


def historical_token(group: list[SourceRow]) -> str:
    identity = ":".join(str(part) for part in group_key(group[0]))
    return str(uuid5(NAMESPACE_URL, f"gastroflow:pedidos-historicos:{identity}"))


def ensure_historical_product(session: Session, products: dict[str, Producto]) -> None:
    if "TERNERA_HISTORICA" in products:
        product = products["TERNERA_HISTORICA"]
        if product.activo:
            raise RuntimeError("TERNERA_HISTORICA debe permanecer inactivo.")
        return
    product = Producto(
        codigo="TERNERA_HISTORICA",
        nombre="Ternera",
        descripcion="Producto histórico vendido por única vez.",
        unidad_venta=UnidadVenta.PIEZA,
        activo=False,
    )
    session.add(product)
    session.flush()
    products[product.codigo] = product


def import_orders(path: Path, *, apply: bool) -> dict[str, object]:
    source_rows = load_source(path)
    groups = group_source(source_rows)
    learned_zones = address_zone_map(source_rows)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    print(
        f"Fuente validada: {len(source_rows)} filas, {len(groups)} pedidos, SHA256 {source_hash}",
        flush=True,
    )

    with Session(engine) as session:
        revision = session.exec(text("SELECT version_num FROM alembic_version")).one()[0]
        if revision != "0011_split_payments_addons":
            raise RuntimeError(f"Revisión Alembic inesperada: {revision}")

        existing_orders = session.exec(select(Pedido).order_by(Pedido.created_at, Pedido.id)).all()
        if len(existing_orders) != 7:
            raise RuntimeError(
                f"Se esperaban 7 pedidos existentes antes de importar y hay {len(existing_orders)}."
            )
        if any(order.fecha_entrega < date(2026, 10, 15) for order in existing_orders):
            raise RuntimeError("Hay pedidos existentes dentro del período histórico del Excel.")
        imported_tokens = {historical_token(group) for group in groups}
        if session.exec(select(func.count()).select_from(Pedido).where(
            Pedido.submission_token.in_(imported_tokens)
        )).one():
            raise RuntimeError("La importación histórica ya fue ejecutada total o parcialmente.")

        categories = {item.codigo: item for item in session.exec(select(Categoria)).all()}
        products = {item.codigo: item for item in session.exec(select(Producto)).all()}
        zones = {normalize_text(item.nombre): item for item in session.exec(select(ZonaEnvio)).all()}
        addons = {item.codigo: item for item in session.exec(select(Agregado)).all()}
        for category_code in set(CATEGORY_CODES.values()):
            if category_code not in categories:
                raise RuntimeError(f"Falta la categoría {category_code}.")
        for product_code in {
            "MUZZA", "ESPECIAL", "CALABRESA", "PREPIZZA", "NAPOLITANA", "FUGAZZETTA", "PANARABE"
        }:
            if product_code not in products:
                raise RuntimeError(f"Falta el producto {product_code}.")
        for zone_name in ("San Miguel de Tucumán", "Alderetes", "Yerba Buena"):
            if normalize_text(zone_name) not in zones:
                raise RuntimeError(f"Falta la zona {zone_name}.")
        if "HUEVO" not in addons:
            raise RuntimeError("Falta el agregado HUEVO.")
        ensure_historical_product(session, products)
        print(
            f"Base validada: {len(existing_orders)} pedidos existentes; catálogo completo.",
            flush=True,
        )

        existing_clients: dict[str, Cliente] = {}
        for client in session.exec(select(Cliente).order_by(Cliente.id)).all():
            key = normalize_text(client.nombre_apellido)
            if key in existing_clients:
                raise RuntimeError(f"Hay clientes existentes duplicados por nombre: {client.nombre_apellido}")
            existing_clients[key] = client

        # Free the definitive code range before inserting the historical orders.
        for order in existing_orders:
            order.codigo = f"TMP-IMPORT-{order.id}"
            session.add(order)
        session.flush()

        created_clients = 0
        imported_item_count = 0
        imported_addon_count = 0
        mixed_order_count = 0
        zone_counts: Counter[str] = Counter()
        historical_orders: list[Pedido] = []

        for order_number, group in enumerate(groups, start=1):
            first = group[0]
            customer_name = corrected_name(first)
            customer_key = normalize_text(customer_name)
            client = existing_clients.get(customer_key)
            created_at = datetime.combine(
                first.delivery_date,
                time(12, 0),
                tzinfo=LOCAL_TZ,
            ) + timedelta(seconds=order_number)
            if client is None:
                client = Cliente(
                    nombre_apellido=customer_name,
                    telefono=None,
                    created_at=created_at,
                    updated_at=created_at,
                )
                session.add(client)
                session.flush()
                existing_clients[customer_key] = client
                created_clients += 1

            shipping_total = sum(
                (row.shipping or Decimal("0.00") for row in group),
                Decimal("0.00"),
            ).quantize(MONEY_QUANT)
            payment_amounts: dict[FormaPago, Decimal] = defaultdict(lambda: Decimal("0.00"))
            for row in group:
                payment_amounts[payment_type(row.payment_text)] += row.line_total + (
                    row.shipping or Decimal("0.00")
                )
            payment_amounts = {
                key: value.quantize(MONEY_QUANT)
                for key, value in payment_amounts.items()
            }
            order_total = (
                sum((row.line_total for row in group), Decimal("0.00")) + shipping_total
            ).quantize(MONEY_QUANT)
            if sum(payment_amounts.values(), Decimal("0.00")) != order_total:
                raise RuntimeError(f"Pagos descuadrados en fila {first.excel_row}.")
            is_mixed = len(payment_amounts) > 1
            if is_mixed:
                mixed_order_count += 1

            zone_name = resolve_group_zone(group, learned_zones)
            zone = zones.get(normalize_text(zone_name)) if zone_name else None
            zone_counts[zone_name or "Sin zona"] += 1
            pedido = Pedido(
                submission_token=historical_token(group),
                codigo=f"PED-{order_number:04d}",
                fecha_entrega=first.delivery_date,
                horario_entrega=None,
                cliente_id=client.id or 0,
                direccion_delivery=(
                    "Retiro en el local"
                    if normalize_text(first.address) == "retiro en el local"
                    else first.address
                ),
                zona_envio_id=zone.id if zone else None,
                costo_envio=shipping_total,
                forma_pago=(FormaPago.MIXTO if is_mixed else next(iter(payment_amounts))),
                estado=EstadoPedido.ENTREGADO,
                observaciones=group_observation(group),
                monto_total=order_total,
                created_at=created_at,
                updated_at=created_at,
            )
            session.add(pedido)
            session.flush()
            historical_orders.append(pedido)

            for row in group:
                product_code, combo_code, quantity = parse_product(row)
                line_subtotal = row.line_total
                has_egg = row.excel_row == 437
                if has_egg:
                    line_subtotal -= Decimal("1000.00")
                if line_subtotal <= 0:
                    raise RuntimeError(f"Subtotal inválido en fila {row.excel_row}.")
                item = PedidoItem(
                    pedido_id=pedido.id or 0,
                    producto_id=products[product_code].id or 0,
                    producto_combo_id=products[combo_code].id if combo_code else None,
                    categoria_venta_id=categories[CATEGORY_CODES[normalize_text(row.category_text)]].id or 0,
                    cantidad=quantity,
                    precio_unitario=(line_subtotal / Decimal(quantity)).quantize(
                        MONEY_QUANT, rounding=ROUND_HALF_UP
                    ),
                    subtotal=line_subtotal,
                    es_especial=False,
                )
                session.add(item)
                session.flush()
                imported_item_count += 1
                if has_egg:
                    egg = addons["HUEVO"]
                    session.add(PedidoItemAgregado(
                        pedido_item_id=item.id or 0,
                        agregado_id=egg.id or 0,
                        agregado_nombre="Huevo",
                        cantidad=1,
                        precio_unitario=Decimal("1000.00"),
                        subtotal=Decimal("1000.00"),
                    ))
                    imported_addon_count += 1

            for form, amount in payment_amounts.items():
                session.add(PedidoPago(
                    pedido_id=pedido.id or 0,
                    forma_pago=form,
                    monto=amount,
                ))
            if order_number % 50 == 0:
                print(f"Preparados {order_number}/{len(groups)} pedidos históricos...", flush=True)

        next_number = len(historical_orders) + 1
        for offset, order in enumerate(existing_orders):
            order.codigo = f"PED-{next_number + offset:04d}"
            session.add(order)
        session.flush()

        final_order_count = session.exec(select(func.count()).select_from(Pedido)).one()
        expected_final_count = len(historical_orders) + len(existing_orders)
        if final_order_count != expected_final_count:
            raise RuntimeError(
                f"Conteo final inesperado: {final_order_count}, esperado {expected_final_count}."
            )
        total_payments = session.exec(select(func.sum(PedidoPago.monto))).one() or Decimal("0.00")
        total_orders = session.exec(select(func.sum(Pedido.monto_total))).one() or Decimal("0.00")
        if total_payments.quantize(MONEY_QUANT) != total_orders.quantize(MONEY_QUANT):
            raise RuntimeError("La suma global de pagos no coincide con la suma global de pedidos.")

        final_code_number = expected_final_count
        if apply:
            session.exec(text(
                "SELECT setval('pedido_codigo_seq', :last_value, true)"
            ), params={"last_value": final_code_number})

        summary = {
            "source_sha256": source_hash,
            "source_rows": len(source_rows),
            "historical_orders": len(historical_orders),
            "historical_items": imported_item_count,
            "historical_clients_created": created_clients,
            "split_payment_orders": mixed_order_count,
            "historical_addons": imported_addon_count,
            "existing_orders_renumbered": len(existing_orders),
            "final_orders": final_order_count,
            "last_code": f"PED-{final_code_number:04d}",
            "next_code": f"PED-{final_code_number + 1:04d}",
            "zone_counts": dict(sorted(zone_counts.items())),
            "historical_total": str(sum(
                (order.monto_total for order in historical_orders), Decimal("0.00")
            ).quantize(MONEY_QUANT)),
            "mode": "APPLY" if apply else "DRY-RUN",
        }
        print("Conciliación final correcta.", flush=True)
        if apply:
            session.commit()
            print("Transacción confirmada.", flush=True)
        else:
            session.rollback()
            print("Ensayo revertido correctamente.", flush=True)
        return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa los pedidos históricos validados.")
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    summary = import_orders(args.file.resolve(), apply=args.apply)
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
