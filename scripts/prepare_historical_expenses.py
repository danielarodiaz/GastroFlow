from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


EXPECTED_HEADERS = ("ID", "Fecha", "Motivo", "Marca", "Cantidad", "Precio", "Lugar")
EXPECTED_SOURCE_ROWS = 230
EXPECTED_TICKET_GROUPS = 109
EXPECTED_FIRST_DATE = date(2026, 6, 12)
EXPECTED_LAST_DATE = date(2026, 10, 2)
EXPECTED_TOTAL = Decimal("2240360.28")
MONEY_QUANT = Decimal("0.01")
QUANTITY_QUANT = Decimal("0.001")


PLACE_CORRECTIONS = {
    "veduleria sofia": "Verduleria Sofia",
    "verduleria sofia": "Verduleria Sofia",
    "elisa": "Distribuidora Elisa",
    "distribuidora elisa": "Distribuidora Elisa",
    "ypf": "YPF",
    "ypf autopista": "YPF",
}

MOTIVE_CORRECTIONS = {
    "pure de tomate": "Pure de Tomate",
}


@dataclass(frozen=True)
class QuantityResult:
    quantity: Decimal
    unit: str
    presentation: str | None = None
    review_required: bool = False
    note: str | None = None


@dataclass(frozen=True)
class PreparedExpenseRow:
    excel_row: int
    source_id: str
    ticket_group: str
    expense_date: date
    place_original: str
    place_normalized: str
    motive_original: str
    motive_normalized: str
    brand_original: str | None
    brand_normalized: str | None
    quantity_original: str | None
    quantity: Decimal
    unit: str
    presentation: str | None
    line_total: Decimal
    unit_price: Decimal
    status: str
    note: str | None


def clean_display(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def normalize_key(value: object) -> str:
    display = clean_display(value)
    return "".join(
        character
        for character in unicodedata.normalize("NFKD", display)
        if not unicodedata.combining(character)
    ).casefold()


def money(value: object) -> Decimal:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(MONEY_QUANT)
    cleaned = clean_display(value).replace("$", "").replace(" ", "")
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    return Decimal(cleaned).quantize(MONEY_QUANT)


def normalize_place(value: object) -> str:
    display = clean_display(value)
    return PLACE_CORRECTIONS.get(normalize_key(display), display)


def normalize_motive(value: object) -> str:
    display = clean_display(value)
    return MOTIVE_CORRECTIONS.get(normalize_key(display), display)


def normalize_brand(value: object) -> str | None:
    display = clean_display(value)
    if not display or normalize_key(display) == "sin marca":
        return None
    return display


def _decimal_number(value: str) -> Decimal:
    return Decimal(value.replace(",", "."))


def parse_quantity(value: object, *, motive: str) -> QuantityResult:
    original = clean_display(value)
    normalized = normalize_key(original).replace(" ", "")
    motive_key = normalize_key(motive)

    if not original:
        if motive_key in {"gomitas de rueda trasera y calibrado", "propina por puesta de camara"}:
            unit = "SERVICIO"
        elif motive_key == "nafta super":
            unit = "CARGA"
        elif motive_key == "verduras":
            unit = "LOTE"
        else:
            unit = "UNIDAD"
        return QuantityResult(
            quantity=Decimal("1.000"),
            unit=unit,
            review_required=False,
            note=f"Cantidad confirmada por el usuario: 1 {unit.lower()}.",
        )

    if normalized == "4ltymedio":
        return QuantityResult(Decimal("4.500"), "L", presentation=original)
    if normalized == "1/2kg":
        return QuantityResult(Decimal("0.500"), "KG", presentation=original)
    if normalized == "tanquelleno":
        return QuantityResult(Decimal("1.000"), "CARGA", presentation=original)
    if normalized == "1rollomediano":
        return QuantityResult(Decimal("1.000"), "ROLLO", presentation=original)
    if normalized == "850lt":
        return QuantityResult(
            Decimal("0.850"),
            "L",
            presentation=original,
            review_required=False,
            note="Cantidad confirmada por el usuario: 850 ml (0,850 L); el origen dice 850lt.",
        )

    package_match = re.fullmatch(r"(mediahorma|hormacompleta)/(\d+(?:[.,]\d+)?)kg", normalized)
    if package_match:
        return QuantityResult(
            _decimal_number(package_match.group(2)).quantize(QUANTITY_QUANT),
            "KG",
            presentation=original.split("/", 1)[0],
        )

    multipack_match = re.fullmatch(r"(\d+)x(\d+(?:[.,]\d+)?)(gr|kg|ml|lt|l)", normalized)
    if multipack_match:
        packages = Decimal(multipack_match.group(1))
        measure = _decimal_number(multipack_match.group(2)) * packages
        source_unit = multipack_match.group(3)
        if source_unit == "gr":
            measure /= Decimal("1000")
            unit = "KG"
        elif source_unit == "kg":
            unit = "KG"
        elif source_unit == "ml":
            measure /= Decimal("1000")
            unit = "L"
        else:
            unit = "L"
        return QuantityResult(measure.quantize(QUANTITY_QUANT), unit, presentation=original)

    measure_match = re.fullmatch(r"(\d+(?:[.,]\d+)?)(kg|gr|ml|lt|l)", normalized)
    if measure_match:
        measure = _decimal_number(measure_match.group(1))
        source_unit = measure_match.group(2)
        if source_unit == "gr":
            measure /= Decimal("1000")
            unit = "KG"
        elif source_unit == "kg":
            unit = "KG"
        elif source_unit == "ml":
            measure /= Decimal("1000")
            unit = "L"
        else:
            unit = "L"
        return QuantityResult(measure.quantize(QUANTITY_QUANT), unit)

    if re.fullmatch(r"\d+(?:[.,]\d+)?", normalized):
        return QuantityResult(_decimal_number(normalized).quantize(QUANTITY_QUANT), "UNIDAD")

    raise RuntimeError(f"Cantidad no contemplada: {original!r}")


def load_and_prepare(path: Path) -> tuple[list[PreparedExpenseRow], list[dict[str, object]]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Gastos" not in workbook.sheetnames:
        raise RuntimeError("El archivo no contiene la hoja Gastos.")
    sheet = workbook["Gastos"]
    headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=7)))
    if headers != EXPECTED_HEADERS:
        raise RuntimeError(f"Encabezados inesperados: {headers!r}")

    source: list[tuple[int, tuple[object, ...]]] = []
    for excel_row, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2
    ):
        if values[1] is None:
            continue
        source.append((excel_row, values))
    workbook.close()

    if len(source) != EXPECTED_SOURCE_ROWS:
        raise RuntimeError(
            f"Se esperaban {EXPECTED_SOURCE_ROWS} filas y se encontraron {len(source)}."
        )

    prepared: list[PreparedExpenseRow] = []
    for excel_row, values in source:
        source_id, raw_date, raw_motive, raw_brand, raw_quantity, raw_price, raw_place = values
        if not isinstance(raw_date, datetime):
            raise RuntimeError(f"Fecha inválida en fila {excel_row}: {raw_date!r}")
        expense_date = raw_date.date()
        place_original = clean_display(raw_place)
        place_normalized = normalize_place(raw_place)
        motive_original = clean_display(raw_motive)
        motive_normalized = normalize_motive(raw_motive)
        brand_original = clean_display(raw_brand) or None
        brand_normalized = normalize_brand(raw_brand)
        quantity_result = parse_quantity(raw_quantity, motive=motive_normalized)
        line_total = money(raw_price)
        if line_total <= 0:
            raise RuntimeError(f"Subtotal inválido en fila {excel_row}: {line_total}")
        unit_price = (line_total / quantity_result.quantity).quantize(
            MONEY_QUANT, rounding=ROUND_HALF_UP
        )
        ticket_group = f"{expense_date.isoformat()}|{normalize_key(place_normalized)}"
        prepared.append(
            PreparedExpenseRow(
                excel_row=excel_row,
                source_id=str(source_id),
                ticket_group=ticket_group,
                expense_date=expense_date,
                place_original=place_original,
                place_normalized=place_normalized,
                motive_original=motive_original,
                motive_normalized=motive_normalized,
                brand_original=brand_original,
                brand_normalized=brand_normalized,
                quantity_original=clean_display(raw_quantity) or None,
                quantity=quantity_result.quantity,
                unit=quantity_result.unit,
                presentation=quantity_result.presentation,
                line_total=line_total,
                unit_price=unit_price,
                status="REVISAR" if quantity_result.review_required else "LISTO",
                note=quantity_result.note,
            )
        )

    dates = [row.expense_date for row in prepared]
    if min(dates) != EXPECTED_FIRST_DATE or max(dates) != EXPECTED_LAST_DATE:
        raise RuntimeError("El rango de fechas no coincide con la fuente validada.")
    total = sum((row.line_total for row in prepared), Decimal("0.00"))
    if total != EXPECTED_TOTAL:
        raise RuntimeError(f"Total inesperado: {total}; esperado: {EXPECTED_TOTAL}.")

    grouped: dict[str, list[PreparedExpenseRow]] = {}
    for row in prepared:
        grouped.setdefault(row.ticket_group, []).append(row)
    if len(grouped) != EXPECTED_TICKET_GROUPS:
        raise RuntimeError(
            f"Se esperaban {EXPECTED_TICKET_GROUPS} tickets y se reconstruyeron {len(grouped)}."
        )

    tickets: list[dict[str, object]] = []
    for number, rows in enumerate(grouped.values(), start=1):
        first = rows[0]
        tickets.append(
            {
                "ticket_number": number,
                "ticket_group": first.ticket_group,
                "expense_date": first.expense_date,
                "place": first.place_normalized,
                "item_count": len(rows),
                "ticket_total": sum(
                    (row.line_total for row in rows), Decimal("0.00")
                ).quantize(MONEY_QUANT),
                "source_rows": ", ".join(str(row.excel_row) for row in rows),
                "review_items": sum(row.status == "REVISAR" for row in rows),
            }
        )
    return prepared, tickets


def _json_value(value: object) -> object:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date):
        return value.isoformat()
    raise TypeError(f"Tipo no serializable: {type(value)!r}")


def build_report(path: Path) -> dict[str, object]:
    rows, tickets = load_and_prepare(path)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    total = sum((row.line_total for row in rows), Decimal("0.00")).quantize(MONEY_QUANT)
    status_counts = Counter(row.status for row in rows)
    report = {
        "summary": {
            "source_file": path.name,
            "source_sha256": source_hash,
            "source_rows": len(rows),
            "ticket_groups": len(tickets),
            "ready_items": status_counts["LISTO"],
            "review_items": status_counts["REVISAR"],
            "items_without_brand": sum(row.brand_normalized is None for row in rows),
            "first_date": min(row.expense_date for row in rows),
            "last_date": max(row.expense_date for row in rows),
            "source_total": total,
            "prepared_total": total,
            "reconciliation_difference": Decimal("0.00"),
        },
        "tickets": tickets,
        "rows": [asdict(row) for row in rows],
    }
    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Valida y prepara la planilla histórica de gastos sin escribir en la base."
    )
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()

    source = args.file.resolve()
    report = build_report(source)
    summary = report["summary"]
    print("Validación histórica de gastos correcta.")
    for key, value in summary.items():
        print(f"{key}: {value}")
    if args.json_output:
        output = args.json_output.resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=_json_value),
            encoding="utf-8",
        )
        print(f"json_output: {output}")


if __name__ == "__main__":
    main()
